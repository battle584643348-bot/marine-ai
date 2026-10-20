"""技能：生成 Excel 工程报告（Punch List / 测试记录表 / 通用清单）。"""
import time
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .base import Skill

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

_SCHEMAS = {
    "punch_list": {
        "title": "Punch List 缺陷清单",
        "headers": ["序号", "系统", "位置", "缺陷描述", "类别", "优先级", "状态", "责任人", "计划完成日期"],
        "widths": [6, 16, 18, 40, 10, 9, 10, 12, 14],
        "keys": ["no", "system", "location", "description", "category", "priority", "status", "responsible", "due_date"],
        "fills": {"P1": "C00000", "P2": "E36C0A", "P3": "BFBFBF"},
    },
    "test_values": {
        "title": "测试记录表 (Test Values)",
        "headers": ["序号", "测试项目", "规范要求", "实测值", "单位", "判定"],
        "widths": [6, 34, 24, 16, 8, 10],
        "keys": ["no", "item", "spec", "measured", "unit", "result"],
        "fills": {"PASS": "C6EFCE", "FAIL": "FFC7CE", "合格": "C6EFCE", "不合格": "FFC7CE"},
    },
    "general": {
        "title": "通用清单",
        "headers": ["序号", "项目", "内容/数值", "备注"],
        "widths": [6, 24, 44, 20],
        "keys": ["no", "item", "value", "remark"],
        "fills": {},
    },
}


class ExcelReportSkill(Skill):
    name = "excel_report"
    title = "Excel 报告生成"
    description = (
        "生成可直接下载的 Excel 工程表格：punch_list 缺陷清单、test_values 测试记录表、general 通用清单。"
        "适合用户要求'生成/导出一份XX清单/记录表(Excel)'的场景。"
    )
    parameters = [
        {"name": "kind", "type": "string", "required": True,
         "enum": ["punch_list", "test_values", "general"],
         "description": "表格类型"},
        {"name": "rows", "type": "array", "required": True,
         "description": "数据行，如 [{\"no\":1,\"system\":\"MCC-01\",...}, ...]，字段按表格类型填写"},
        {"name": "title", "type": "string", "required": False, "description": "工作表/标题名，可选"},
    ]

    async def run(self, ctx, args: Dict[str, Any]) -> Dict[str, Any]:
        kind = str(args.get("kind", "")).strip()
        schema = _SCHEMAS.get(kind)
        if not schema:
            return {"type": "text", "content": f"不支持的表格类型：{kind}。可选：{', '.join(_SCHEMAS)}"}
        rows = args.get("rows")
        if not isinstance(rows, list):
            return {"type": "text", "content": "参数缺失或格式错误：rows 需为数组，每行含表格对应字段。"}
        title = str(args.get("title") or schema["title"])

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        for ci, (h, w) in enumerate(zip(schema["headers"], schema["widths"]), start=1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = _CENTER
            cell.border = _BORDER
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2"

        for ri, row in enumerate(rows, start=2):
            for ci, key in enumerate(schema["keys"], start=1):
                val = row.get(key, "") if isinstance(row, dict) else (row[ci - 1] if ci - 1 < len(row) else "")
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.border = _BORDER
                cell.alignment = _CENTER if ci in (1, 3, 4, 5, 6) else _LEFT
                fill_key = schema["fills"].get(str(val))
                if fill_key:
                    cell.fill = PatternFill("solid", fgColor=fill_key)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(schema['headers']))}{ws.max_row}"

        generated_dir = Path(ctx.paths["generated"])
        generated_dir.mkdir(parents=True, exist_ok=True)
        filename = f"{int(time.time()*1000)}_{kind}.xlsx"
        path = generated_dir / filename
        wb.save(str(path))
        return {"type": "file", "content": filename,
                "meta": {"path": str(path), "filename": filename, "title": title, "rows": len(rows)}}
