# -*- coding: utf-8 -*-
"""
技能：单线图按系统自动标注（照明版 + 开关柜版，一个技能两种模式）

功能（由用户脚本 mark_lighting_diagram.py / mark_sld_diagram.py 合并封装）：
  mode="lighting"：照明/配电单线图 —— 识别回路(CKT/LOOP)，按系统框出配电板与
                   回路，加 <系统号>-<回路号> / LOOP-<回路号> 可编辑标签；
  mode="sld"     ：开关柜单线图   —— 识别 P 动力电缆馈线与尾设备标签，按子系统
                   框住"电缆+下游设备"，标注母线(BUSBAR)，栅格化图纸自动 OCR，
                   并将识别结果回填到清册副本(新增 SLD识别结果 工作表)。
  所有框/文本框均为 PDF 注释(Annotation)，可在福昕/PDF-XChange 中选中、拖动、
  双击编辑，不烘焙进图纸；同时导出"未匹配电缆清单.xlsx"。

依赖：pymupdf、openpyxl（OCR 可选：rapidocr_onnxruntime）
"""
import os
import re
import time
from collections import Counter, OrderedDict
from pathlib import Path
from typing import Any, Dict, List

import pymupdf  # PyMuPDF
import openpyxl

from .base import Skill

# ===========================================================================
# ① lighting 模式（照明/配电单线图）
# ===========================================================================
L_KNOWN_COLORS = {
    "80-ELE-80-07": (0.05, 0.35, 0.85),   # 正常照明 - 蓝
    "80-ELE-90-08": (0.85, 0.12, 0.12),   # 应急照明 - 红
}
L_PALETTE = [
    (0.05, 0.35, 0.85), (0.85, 0.12, 0.12), (0.10, 0.55, 0.25), (0.90, 0.55, 0.05),
    (0.50, 0.20, 0.70), (0.05, 0.60, 0.60), (0.75, 0.35, 0.55), (0.40, 0.40, 0.10),
]
L_LINE_WIDTH = 2.0
L_DB_TAG_FS = 9.0
L_LOOP_TAG_FS = 8.0
L_PANEL_LEFT = 55.0
L_PANEL_TOP_ABS = 110.0
L_PANEL_BOT_PAD = 225.0
L_PANEL_RIGHT_PAD = 65.0
L_LOOP_LEFT_PAD = 0.0
L_LOOP_TOP_PAD = 0.0
L_LOOP_RIGHT_PAD = 12.0
L_LOOP_BOT_PAD = 12.0
L_DB_TAG_X = 480.0
L_CABLE_RE = re.compile(r"^(P|I|C)-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+\d{2}-\d{2}$")
L_DB_RE = re.compile(r"^DB\d+[A-Z]?\d+$")
L_LOOP_RE = re.compile(r"-[A-Z](\d{2})\d{2}-")


def lighten(c, f=0.85):
    return tuple(min(1.0, v + f * (1.0 - v)) for v in c)


def norm(s: str) -> str:
    """归一化：软连字符->普通连字符，转大写，去首尾空白与标点。"""
    return s.replace("\u00ad", "-").upper().strip().rstrip(",.;:")


def l_recover_cable(tok: str) -> str:
    """修复 PDF 字体 ToUnicode 缺陷：数字 0 被错抽成字母 Q。"""
    if L_CABLE_RE.match(tok):
        return tok
    repaired = tok.replace("Q", "0")
    if L_CABLE_RE.match(repaired):
        return repaired
    return tok


def l_add_freetext(page, doc, rect, text, fs, col, border, fill, align=0):
    """福昕风格可编辑文本框（FreeText，可双击编辑/拖动）。中文须 china-s。"""
    a = page.add_freetext_annot(
        rect, text, fontname="china-s", fontsize=fs,
        text_color=col, fill_color=fill,
        border_color=border, align=align, richtext=True,
    )
    try:
        a.set_border(width=L_LINE_WIDTH, style="S")
    except Exception:
        pass
    a.update()
    try:
        da = "%.3f %.3f %.3f rg /china-s %.1f Tf" % (col[0], col[1], col[2], fs)
        doc.xref_set_key(a.xref, "DA", "(" + da + ")")
    except Exception:
        pass
    return a


def l_add_rectframe(page, rect, col, width=L_LINE_WIDTH):
    a = page.add_rect_annot(rect)
    a.set_colors(stroke=col)
    try:
        a.set_border(width=width, style="S")
    except Exception:
        pass
    a.update()
    return a


def l_find_db_token(page):
    for w in page.get_text("words"):
        t = w[4].replace("\u00ad", "-").upper()
        if L_DB_RE.match(t) and w[0] < 150 and w[1] < 110:
            return t
    return None


def l_panel_bounds(words, pr):
    mc = [w for w in words if any(k in w[4].upper() for k in ("MCB", "MCCB", "BUS"))]
    if not mc:
        return None
    px1 = max(w[2] for w in mc) + L_PANEL_RIGHT_PAD
    py0 = L_PANEL_TOP_ABS if L_PANEL_TOP_ABS > 0 else max(min(w[1] for w in mc) - 18, 95)
    py1 = min(max(w[3] for w in mc) + L_PANEL_BOT_PAD, pr.y1 - 95)
    return (L_PANEL_LEFT, py0, px1, py1)


def l_build_lookup(schedule_path: str, sheet_name: str | None):
    wb = openpyxl.load_workbook(schedule_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    lookup, panel_sys, panel_desc, descs = {}, {}, {}, {}
    for row in ws.iter_rows(values_only=True):
        if len(row) < 36 or not row[4]:
            continue
        cable = norm(str(row[4]))
        sysno = str(row[34]).strip() if row[34] else ""
        desc = str(row[35]).strip() if row[35] else ""
        if not sysno:
            continue
        if cable not in lookup:
            lookup[cable] = sysno
        parts = cable.split("-")
        if len(parts) >= 2 and parts[1].startswith("DB"):
            panel = parts[1]
            if panel not in panel_sys or len(sysno) > len(panel_sys[panel]):
                panel_sys[panel] = sysno
                panel_desc[panel] = desc
        if sysno and desc and sysno not in descs:
            descs[sysno] = desc
    wb.close()
    return lookup, panel_sys, panel_desc, descs


def l_annotate(src_pdf, out_pdf, lookup, panel_sys, panel_desc, descs):
    doc = pymupdf.open(src_pdf)
    color_of = OrderedDict()
    loop_color_idx = [0]

    def color_for(sysno):
        if sysno in L_KNOWN_COLORS:
            return L_KNOWN_COLORS[sysno]
        if sysno not in color_of:
            color_of[sysno] = L_PALETTE[len(color_of) % len(L_PALETTE)]
        return color_of[sysno]

    def next_loop_color():
        c = L_PALETTE[loop_color_idx[0] % len(L_PALETTE)]
        loop_color_idx[0] += 1
        return c

    frames = freetexts = 0
    unmatched = []
    seen = set()
    loop_pages = 0

    for pno in range(doc.page_count):
        page = doc[pno]; pr = page.rect
        if not any(f[4] == "china-s" for f in page.get_fonts()):
            page.insert_font(fontname="china-s")
        db = l_find_db_token(page)
        words = page.get_text("words")

        cables_raw = []
        for w in words:
            c = l_recover_cable(norm(w[4]))
            if L_CABLE_RE.match(c):
                cables_raw.append((w[0], w[1], w[2], w[3], c))
                seen.add(c)

        ckts = []
        for w in words:
            if w[4].replace("\u00ad", "-").upper() == "CKT":
                cx, cy = w[2], w[1]
                num = None
                for w2 in words:
                    if w2[0] >= cx and abs(w2[1] - cy) < 9:
                        t2 = w2[4].replace("\u00ad", "").replace("Q", "0")
                        if re.match(r"^\d{1,3}$", t2):
                            num = t2
                            break
                if num and int(num) != 0:
                    ckts.append((cx, cy, f"{int(num):02d}"))

        loop_list = []
        for (cx, cy, num) in ckts:
            g = [c for c in cables_raw if L_LOOP_RE.search(c[4]) and L_LOOP_RE.search(c[4]).group(1) == num]
            if not g:
                near = [c[4] for c in cables_raw if abs((c[1] + c[3]) / 2 - cy) < 55]
                for cnum in near:
                    unmatched.append((pno + 1, cnum, "回路段号格式未识别"))
                continue
            first = min(g, key=lambda t: t[0])
            fc = first[4]
            s = lookup.get(fc)
            if not s:
                parts = fc.split("-")
                if len(parts) >= 2 and parts[1].startswith("DB"):
                    s = panel_sys.get(parts[1])
            if not s:
                unmatched.append((pno + 1, fc, "清册中无此电缆"))
            loop_list.append((cx, cy, num, g, s))

        resolved = [t[4] for t in loop_list if t[4]]
        if not loop_list and db is None:
            continue
        sysdb = Counter(resolved).most_common(1)[0][0] if resolved else None
        if sysdb is None:
            continue
        dbcol = color_for(sysdb)
        loop_pages += 1

        pb = l_panel_bounds(words, pr)
        if pb:
            l_add_rectframe(page, pymupdf.Rect(*pb), dbcol)
            frames += 1
        px_right = pb[2] + 10 if pb else 190.0

        obstacles = [pymupdf.Rect(w[0] - 15, w[1] - 30, w[2] + 15, w[3] + 15)
                     for w in words if norm(w[4]).startswith("SB") and "SWITCHBOARD" not in norm(w[4])]

        desc = panel_desc.get(db) or descs.get(sysdb, "")
        dbw = min(pr.x1 - 10 - L_DB_TAG_X, 700)
        db_rect = pymupdf.Rect(L_DB_TAG_X, 96, L_DB_TAG_X + dbw,
                               96 + 24 + (16 if desc else 0))
        l_add_freetext(page, doc, db_rect, sysdb + ("\n" + desc if desc else ""),
                       L_DB_TAG_FS, dbcol, dbcol, lighten(dbcol), 0)
        freetexts += 1

        for (cx, cy, num, g, loop_sys) in loop_list:
            loopcol = next_loop_color()
            xs0 = min(c[0] for c in g); xs1 = max(c[2] for c in g)
            ys0 = min(c[1] for c in g); ys1 = max(c[3] for c in g)
            for w in words:
                wy = (w[1] + w[3]) / 2
                if ys0 - 15 <= wy <= ys1 + 15 and w[0] >= xs0 - 20:
                    xs1 = max(xs1, w[2]); ys1 = max(ys1, w[3]); ys0 = min(ys0, w[1])
            x0f = max(xs0 - L_LOOP_LEFT_PAD, px_right)
            y0f = max(ys0 - L_LOOP_TOP_PAD, 40)
            x1f = min(xs1 + L_LOOP_RIGHT_PAD, pr.x1 - 30)
            y1f = min(ys1 + L_LOOP_BOT_PAD, pr.y1 - 90)
            l_add_rectframe(page, pymupdf.Rect(x0f, y0f, x1f, y1f), loopcol)
            frames += 1

            fs = L_LOOP_TAG_FS
            l1 = f"{loop_sys}-{num}" if loop_sys else f"未匹配-{num}"
            l2 = f"LOOP-{num}"
            w1 = pymupdf.get_text_length(l1, fontname="china-s", fontsize=fs)
            bw = max(w1, pymupdf.get_text_length(l2, fontname="china-s", fontsize=fs)) + 28
            bh = fs * 5.2 + 4
            ty1 = max(y0f - 6, 100); ty0 = max(ty1 - bh, 88)
            tagr = pymupdf.Rect(x0f + 6, ty0, x0f + 6 + bw, ty1)
            for _ in range(3):
                hit = next((ob for ob in obstacles if tagr.intersects(ob)), None)
                if not hit:
                    break
                ty1 = max(hit.y0 - 4, 88); ty0 = max(ty1 - bh, 88)
                tagr = pymupdf.Rect(x0f + 6, ty0, x0f + 6 + bw, ty1)
            l_add_freetext(page, doc, tagr, l1 + "\n" + l2, fs, loopcol, loopcol, lighten(loopcol), 0)
            freetexts += 1

        print(f"  page {pno+1:>3}: panel={db} system={sysdb} loops={len(loop_list)}")

    doc.save(out_pdf, garbage=4, deflate=True)
    doc.close()
    print(f"=== 照明标注完成: 页{loop_pages} 框{frames} 文本框{freetexts} 未匹配{len(unmatched)} ===")
    return unmatched, seen


def l_write_unmatched_excel(path, drawing_unmatched, schedule_not_seen, lookup, descs):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    hdr_font = Font(color="FFFFFF", bold=True)
    ws1 = wb.active
    ws1.title = "图纸未匹配电缆"
    ws1.append(["序号", "图纸页码", "电缆号", "说明"])
    for i, (pg, c, note) in enumerate(drawing_unmatched, 1):
        ws1.append([i, pg, c, note])
    for col in ("A", "B", "C", "D"):
        ws1[col + "1"].fill = hdr_fill; ws1[col + "1"].font = hdr_font
        ws1[col + "1"].alignment = Alignment(horizontal="center")
    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 10
    ws1.column_dimensions["C"].width = 34
    ws1.column_dimensions["D"].width = 22
    ws1.freeze_panes = "A2"
    ws2 = wb.create_sheet("清册未上图电缆")
    ws2.append(["序号", "电缆号", "系统号", "系统描述"])
    for i, c in enumerate(schedule_not_seen, 1):
        sysno = lookup.get(c, "")
        ws2.append([i, c, sysno, descs.get(sysno, "")])
    for col in ("A", "B", "C", "D"):
        ws2[col + "1"].fill = hdr_fill; ws2[col + "1"].font = hdr_font
        ws2[col + "1"].alignment = Alignment(horizontal="center")
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 34
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["D"].width = 40
    ws2.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        base, ext = os.path.splitext(path)
        path = f"{base}_{time.strftime('%H%M%S')}{ext}"
        wb.save(path)
    return len(drawing_unmatched), len(schedule_not_seen), path


# ===========================================================================
# ② sld 模式（开关柜单线图）
# ===========================================================================
S_PALETTE = [
    (0.85, 0.12, 0.12), (0.05, 0.35, 0.85), (0.10, 0.55, 0.25), (0.90, 0.55, 0.05),
    (0.50, 0.20, 0.70), (0.05, 0.60, 0.60), (0.75, 0.35, 0.55), (0.40, 0.40, 0.10),
]
S_LINE_WIDTH = 2.0
S_TAG_FS = 8.5
S_BUS_FS = 9.5
S_EQUIP_PAD_Y = 6.0
S_TAG_GAP_Y = 8.0
S_SKIP_TOTAG = {"HOLD", "TBC", "TBD", "NA", "N/A", "-", ""}
S_CABLE_RE = re.compile(r'^[PIC]-[A-Z0-9][A-Z0-9-]*\d{2}(?:~\d{2})?$')
S_SB_RE = re.compile(r'^SB\d+[A-Z]?\d*$')
S_POWER_RE = re.compile(r'^P-[A-Z0-9][A-Z0-9-]*\d{2}(?:~\d{2})?$')
S_EQUIP_TAG_RE = re.compile(r'^(?:RCU-)?\d{2}-[A-Z]{2,3}-\d{4}[A-Z]?$')
S_BOARD_RE = re.compile(r'^(DB|SB|ATB|MSB|PDB|ECB|SW)\d')

_ocr_engine = None


def s_get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            print(f"  [警告] RapidOCR 不可用，栅格化图纸将无法标注: {e}")
            _ocr_engine = False
    return _ocr_engine or None


def s_ocr_page(page, dpi=200):
    import tempfile
    ocr = s_get_ocr()
    if ocr is None:
        return []
    pix = page.get_pixmap(dpi=dpi)
    img = os.path.join(tempfile.mkdtemp(), "ocr.png")
    pix.save(img)
    try:
        result, _ = ocr(img)
    finally:
        try:
            os.unlink(img)
        except OSError:
            pass
    if not result:
        return []
    sx = page.rect.width / pix.width
    sy = page.rect.height / pix.height
    words = []
    for box, text, conf in result:
        xs = [p[0] for p in box]; ys = [p[1] for p in box]
        bx0, bx1 = min(xs) * sx, max(xs) * sx
        by0, by1 = min(ys) * sy, max(ys) * sy
        for tok in text.split():
            words.append((bx0, by0, bx1, by1, tok))
    return words


def s_build_lookup(schedule_path: str, sheet_name: str):
    wb = openpyxl.load_workbook(schedule_path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    def find_col(*keys):
        for i, h in enumerate(header):
            ht = str(h).upper() if h else ""
            if any(k.upper() in ht for k in keys):
                return i
        return None

    C_CABLE = find_col("CABLE TAG", "电缆编号", "电缆号")
    C_TOTAG = find_col("TO EQUIP", "尾设备")
    C_SYSNO = find_col("SYSTEM NO", "系统号")
    C_SYSDESC = find_col("SYSTEM DESCRIPTION", "系统描述")
    if C_CABLE is None or C_SYSNO is None:
        raise ValueError(f"清册表头识别失败：未找到电缆号/系统号列。表头={header}")
    print(f"  列定位 -> 电缆号={C_CABLE} 尾设备tag={C_TOTAG} 系统号={C_SYSNO} 系统描述={C_SYSDESC}")
    cable_lu = {}
    equip_lu = {}
    sysdesc = {}
    for row in rows:
        if not row[C_CABLE]:
            continue
        cable = norm(str(row[C_CABLE]))
        sysno = str(row[C_SYSNO]).strip() if row[C_SYSNO] else ""
        desc = str(row[C_SYSDESC]).strip() if C_SYSDESC is not None and row[C_SYSDESC] else ""
        to_tag = norm(str(row[C_TOTAG])) if C_TOTAG is not None and row[C_TOTAG] else ""
        if not sysno:
            continue
        if cable not in cable_lu:
            cable_lu[cable] = (sysno, desc)
        if to_tag and to_tag not in S_SKIP_TOTAG and to_tag not in equip_lu:
            equip_lu[to_tag] = (sysno, desc)
        if sysno not in sysdesc and desc:
            sysdesc[sysno] = desc
    wb.close()
    return cable_lu, equip_lu, sysdesc


def s_expand_cable(c: str):
    m = re.match(r'^(.*-)(\d{2})~(\d{2})$', c)
    if not m:
        return [c]
    head, a, b = m.group(1), int(m.group(2)), int(m.group(3))
    return [f"{head}{i:02d}" for i in range(a, b + 1)]


def s_add_freetext(page, doc, rect, text, fs, col, border, fill):
    a = page.add_freetext_annot(rect, text, fontname="helv", fontsize=fs,
                                text_color=col, fill_color=fill,
                                align=0, richtext=False)
    try:
        a.set_border(width=S_LINE_WIDTH, style="S")
        a.set_colors(stroke=border)
    except Exception:
        pass
    a.update()
    da = "/Helv %.1f Tf %.3f %.3f %.3f rg" % (fs, col[0], col[1], col[2])
    doc.xref_set_key(a.xref, "DA", "(" + da + ")")
    return a


def s_add_rectframe(page, rect, col):
    a = page.add_rect_annot(rect)
    a.set_colors(stroke=col)
    try:
        a.set_border(width=S_LINE_WIDTH, style="S")
    except Exception:
        pass
    a.update()
    return a


def s_annotate(src_pdf, out_pdf, cable_lu, equip_lu, sysdesc):
    doc = pymupdf.open(src_pdf)
    color_of = OrderedDict()

    def color_for(sysno):
        if sysno not in color_of:
            color_of[sysno] = S_PALETTE[len(color_of) % len(S_PALETTE)]
        return color_of[sysno]

    sb_tally = Counter()
    for _pg in doc:
        for w in _pg.get_text("words"):
            t = norm(w[4])
            if S_SB_RE.match(t) and t in equip_lu:
                sysno, dsc = equip_lu[t]
                if "SWITCHBOARD" in (dsc or sysdesc.get(sysno, "")).upper():
                    sb_tally[sysno] += 1
    sw_sys = sb_tally.most_common(1)[0][0] if sb_tally else None
    sw_desc = sysdesc.get(sw_sys, "") if sw_sys else ""

    frames = texts = 0
    unmatched = []
    recognized = []

    for pno in range(doc.page_count):
        page = doc[pno]
        pr = page.rect
        for old in list(page.annots() or []):
            page.delete_annot(old)

        W = [(w[0], w[1], w[2], w[3], norm(w[4])) for w in page.get_text("words")]

        if not any(S_POWER_RE.match(w[4]) for w in W):
            _ocr_w = s_ocr_page(page)
            if _ocr_w:
                if pno == 0:
                    print("  [OCR] 检测到栅格化图纸，启用 OCR 识别")
                W = [(x0, y0, x1, y1, norm(t)) for (x0, y0, x1, y1, t) in _ocr_w]
        pcables = []
        seen_c = set()
        for (x0, y0, x1, y1, t) in W:
            if not S_POWER_RE.match(t):
                continue
            key = (round((x0 + x1) / 2 / 20), round(y0 / 20))
            if key in seen_c:
                continue
            seen_c.add(key)
            is_tba = ('TBA' in t) or ('TBC' in t)
            sysno = desc = None
            if not is_tba:
                for cand in s_expand_cable(t):
                    if cand in cable_lu:
                        sysno, desc = cable_lu[cand]
                        break
            pcables.append(((x0 + x1) / 2, x0, y0, x1, y1, t, sysno, desc, is_tba))
        pcables.sort(key=lambda c: c[0])

        if not pcables:
            continue

        bus = None
        sw_tag = None
        bsys = bdesc = None
        bus_words = [w for w in W
                     if (("BUSBAR" in w[4] or S_SB_RE.match(w[4])) and w[1] < 220)]

        page_board_tag = page_board_sys = page_board_desc = None
        _bc = []
        for (wx0, wy0, wx1, wy1, wt) in W:
            if wy0 > 115:
                continue
            if S_CABLE_RE.match(wt):
                continue
            if wt in equip_lu:
                _sys, _dsc = equip_lu[wt]
                score = 0
                if S_BOARD_RE.match(wt):
                    score += 2
                if "BOARD" in (_dsc or sysdesc.get(_sys, "")).upper():
                    score += 1
                _bc.append((-score, wy0, wx0, wt, _sys, _dsc))
        if _bc:
            _bc.sort()
            _, _, _, page_board_tag, page_board_sys, page_board_desc = _bc[0]

        has_bus = bool(bus_words)
        if not has_bus:
            for d in page.get_drawings():
                for it in d["items"]:
                    if it[0] == "l":
                        p1, p2 = it[1], it[2]
                        if (abs(p1.y - p2.y) < 2 and 100 < p1.y < 220
                                and abs(p2.x - p1.x) > pr.width * 0.4):
                            has_bus = True; break
                    elif it[0] == "re":
                        r = it[1]
                        if 100 < r.y0 < 220 and r.width > pr.width * 0.4:
                            has_bus = True; break
                if has_bus:
                    break
        if not has_bus:
            continue

        if bus_words:
            sw_tag = next((w[4] for w in bus_words if w[4].startswith("SB")), None)
            bsys, bdesc = equip_lu.get(sw_tag, (None, "")) if sw_tag else (None, "")
            if bsys is None:
                c = Counter(c[6] for c in pcables if c[6])
                if c:
                    bsys = c.most_common(1)[0][0]
                    bdesc = sysdesc.get(bsys, "")
            if bsys:
                bx0 = min(w[0] for w in bus_words); by0 = min(w[1] for w in bus_words)
                bx1 = max(w[2] for w in bus_words); by1 = max(w[3] for w in bus_words)
                bus = (bx0, by0, bx1, by1, bsys, bdesc)

        feeder_count = 0
        feeds = []
        n_pc = len(pcables)
        for i, (cx, x0, y0, x1, y1, raw, sysno, desc, is_tba) in enumerate(pcables):
            lane_l = max(45, (pcables[i - 1][0] + cx) / 2 if i > 0 else 45)
            lane_r = min(pr.x1 - 45, (cx + pcables[i + 1][0]) / 2 if i < n_pc - 1 else pr.x1 - 45)
            min_x, max_x, ey_bot = x0, x1, y1 + 12
            equip_tag = ""
            desc_words = []
            for (wx0, wy0, wx1, wy1, wt) in W:
                if S_CABLE_RE.match(wt):
                    continue
                wcx = (wx0 + wx1) / 2
                if lane_l <= wcx <= lane_r and y1 - 15 <= wy0 <= y1 + 170:
                    min_x = min(min_x, wx0); max_x = max(max_x, wx1)
                    ey_bot = max(ey_bot, wy1)
                    if not equip_tag and S_EQUIP_TAG_RE.match(wt):
                        equip_tag = wt
                        continue
                    if wt.upper().startswith("LOC"):
                        continue
                    desc_words.append((round(wy0, 1), wx0, wt))
            desc_words.sort()
            drawn_desc = " ".join(w for _, _, w in desc_words)
            bus_tag = sw_tag or page_board_tag or bsys or sw_sys or ""
            bus_desc = bdesc or page_board_desc or sw_desc or ""
            status = "TBA" if is_tba else ("已匹配" if sysno else "未匹配")
            equip_desc = desc or drawn_desc
            recognized.append((pno + 1, raw, bus_tag, bus_desc, equip_tag, equip_desc, status))
            PAD = 6
            left = max(min_x - PAD, lane_l)
            right = min(max_x + PAD, lane_r)
            top, bottom = y0 - 35, ey_bot + S_EQUIP_PAD_Y
            base = re.sub(r'-\d{2}(?:~\d{2})?$', '', raw)
            if is_tba:
                feeds.append((left, top, right, bottom, ey_bot, (0, 0, 0), "TBA Wait for Update", base))
            elif sysno:
                col = color_for(sysno)
                label = sysno + "\n" + (desc or sysdesc.get(sysno, ""))
                feeds.append((left, top, right, bottom, ey_bot, col, label, base))
            else:
                unmatched.append((pno + 1, raw, "清册中无此动力电缆"))
                feeds.append((left, top, right, bottom, ey_bot, (0, 0, 0), "TBA Wait for Update", base))

        feeds.sort(key=lambda r: r[0])
        GAP = 40
        clusters = []
        for f in feeds:
            if clusters and f[0] - clusters[-1][2] < GAP and f[7] == clusters[-1][7]:
                c = clusters[-1]
                clusters[-1] = (min(c[0], f[0]), min(c[1], f[1]), max(c[2], f[2]),
                                max(c[3], f[3]), max(c[4], f[4]), c[5], c[6], c[7])
            else:
                clusters.append(f)
        for (l, t, rr, b, ey_bot, col, label, base) in clusters:
            s_add_rectframe(page, pymupdf.Rect(l, t, rr, b), col); frames += 1
            th = 30 if label == "TBA Wait for Update" else 44
            tbx = pymupdf.Rect(l, b + S_TAG_GAP_Y - 5, rr, b + S_TAG_GAP_Y + th)
            s_add_freetext(page, doc, tbx, label, S_TAG_FS, col, col, lighten(col, 0.82)); texts += 1
            feeder_count += 1

        if bus:
            bx0, by0, bx1, by1, bsys, bdesc = bus
            bcol = color_for(bsys)
            s_add_rectframe(page, pymupdf.Rect(45, by0 - 6, pr.x1 - 45, by1 + 30), bcol)
            frames += 1
            tbx = pymupdf.Rect(45, by0 - 58, 345, by0 - 8)
            s_add_freetext(page, doc, tbx, bsys + "\n" + (bdesc or sysdesc.get(bsys, "")),
                           S_BUS_FS, bcol, bcol, lighten(bcol, 0.82)); texts += 1
        elif pcables and (page_board_sys or sw_sys):
            bsys = page_board_sys or sw_sys
            bdesc = page_board_desc or sw_desc
            bcol = color_for(bsys)
            s_add_rectframe(page, pymupdf.Rect(45, 108, pr.x1 - 45, 205), bcol)
            frames += 1
            tbx = pymupdf.Rect(45, 50, 345, 100)
            s_add_freetext(page, doc, tbx, bsys + "\n" + bdesc,
                           S_BUS_FS, bcol, bcol, lighten(bcol, 0.82)); texts += 1
            bus = (0, 0, 0, 0, bsys, bdesc)

        if pcables or bus:
            print(f"  page {pno+1:>3}: feeders={feeder_count} pcables={len(pcables)} "
                  f"bus={'Y' if bus else '-'} systems={list(color_of.keys())[-4:]}")

    doc.save(out_pdf, garbage=4, deflate=True)
    doc.close()
    print(f"  -> {os.path.basename(out_pdf)}  框{frames} 文本框{texts} 未匹配{len(unmatched)} "
          f"({os.path.getsize(out_pdf)/1024:.0f} KB)")
    return unmatched, recognized


def s_safe_save_workbook(wb, path):
    try:
        wb.save(path)
        return path
    except PermissionError:
        root, ext = os.path.splitext(path)
        alt = f"{root}_new{ext}"
        wb.save(alt)
        print(f"  [提示] {os.path.basename(path)} 正被占用，已另存为 {os.path.basename(alt)}")
        return alt


def s_write_unmatched_excel(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "未匹配电缆"
    ws.append(["序号", "图纸页码", "电缆号", "说明"])
    hdr = PatternFill("solid", fgColor="C00000")
    for c in "ABCD":
        ws[c + "1"].fill = hdr; ws[c + "1"].font = Font(color="FFFFFF", bold=True)
        ws[c + "1"].alignment = Alignment(horizontal="center")
    for i, (pg, cable, note) in enumerate(rows, 1):
        ws.append([i, pg, cable, note])
    for col, w in zip("ABCD", [8, 10, 34, 22]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return s_safe_save_workbook(wb, path)


def s_drawing_no_from_path(pdf_path):
    m = re.search(r'(SD-\d+)', os.path.basename(pdf_path))
    return m.group(1) if m else os.path.basename(pdf_path)


def s_append_recognized_to_schedule(schedule_path, out_path, rows):
    import shutil
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    shutil.copyfile(schedule_path, out_path)
    wb = load_workbook(out_path)
    name = "SLD识别结果"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(["图号", "页码", "母线tag", "母线描述", "设备tag", "设备描述", "电缆号", "状态"])
    hdr = PatternFill("solid", fgColor="1F4E79")
    for c in "ABCDEFGH":
        ws[c + "1"].fill = hdr; ws[c + "1"].font = Font(color="FFFFFF", bold=True)
        ws[c + "1"].alignment = Alignment(horizontal="center")
    fill_tba = PatternFill("solid", fgColor="D9D9D9")
    for (dno, pg, cab, bt, bd, et, ed, st) in rows:
        ws.append([dno, pg, bt, bd, et, ed, cab, st])
        r = ws.max_row
        if st == "TBA":
            for c in "ABCDEFGH":
                ws[c + str(r)].fill = fill_tba
    for col, w in zip("ABCDEFGH", [12, 8, 16, 30, 22, 36, 28, 10]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return s_safe_save_workbook(wb, out_path)


# ===========================================================================
# ③ 技能入口（一个技能，两种模式）
# ===========================================================================
DEFAULT_SCHEDULE = r"C:\Users\zhang\Desktop\总的电缆清册_系统划分补全_20261012.xlsx"
DEFAULT_SHEET = "cable schedule"


class DiagramAnnotateSkill(Skill):
    name = "diagram_annotate"
    title = "单线图系统标注（照明 / 开关柜）"
    description = (
        "读取电缆清册（含系统号/系统描述），在单线图 PDF 上按系统自动添加可编辑注释框与标签："
        "mode=lighting 标注照明/配电单线图（配电板框、回路框、<系统号>-<回路号> 标签）；"
        "mode=sld 标注开关柜单线图（馈线+设备框、母线框、子系统标签，栅格化图纸自动 OCR，"
        "识别结果回填清册副本）。并导出未匹配电缆清单 Excel。"
        "适合用户要求'给照明/单线图/SLD PDF 按系统加标注/画框'的场景。"
    )
    parameters = [
        {"name": "mode", "type": "string", "required": True,
         "enum": ["lighting", "sld"],
         "description": "标注类型：lighting=照明/配电单线图；sld=开关柜单线图"},
        {"name": "input_pdf", "type": "string", "required": True,
         "description": "待标注单线图 PDF 的完整路径"},
        {"name": "schedule", "type": "string", "required": False,
         "description": "电缆清册 XLSX 完整路径（含系统号/系统描述列）"},
        {"name": "sheet", "type": "string", "required": False,
         "description": "清册中电缆表的 sheet 名（默认 cable schedule）"},
        {"name": "output_dir", "type": "string", "required": False,
         "description": "输出目录（默认系统 generated/diagram_mark 目录）"},
    ]

    async def run(self, ctx, args: Dict[str, Any]) -> Dict[str, Any]:
        mode = str(args.get("mode", "")).strip().lower()
        if mode not in ("lighting", "sld"):
            return {"type": "text",
                    "content": "参数缺失或无效：mode 必须为 lighting（照明单线图）或 sld（开关柜单线图）。"}
        input_pdf = str(args.get("input_pdf", "")).strip().strip('"')
        if not input_pdf or not os.path.exists(input_pdf):
            return {"type": "text", "content": "参数缺失或路径无效：input_pdf 需为存在的单线图 PDF 路径。"}
        schedule = str(args.get("schedule") or DEFAULT_SCHEDULE).strip().strip('"')
        if not os.path.exists(schedule):
            return {"type": "text", "content": f"找不到电缆清册：{schedule}（可用参数 schedule 指定 xlsx 路径）。"}
        sheet = str(args.get("sheet") or DEFAULT_SHEET).strip()

        out_root = Path(ctx.paths["generated"]) / "diagram_mark"
        if args.get("output_dir"):
            out_root = Path(str(args["output_dir"]).strip().strip('"'))
        out_root.mkdir(parents=True, exist_ok=True)

        ts = int(time.time() * 1000)
        stem = os.path.splitext(os.path.basename(input_pdf))[0]
        out_pdf = out_root / f"{ts}_{stem}_系统标注.pdf"

        try:
            if mode == "lighting":
                lookup, panel_sys, panel_desc, descs = l_build_lookup(schedule, sheet)
                print(f"读取电缆清册: {schedule}（sheet={sheet}）")
                print(f"  精确查找表 {len(lookup)} 条电缆；配电板映射 {len(panel_sys)} 个；系统描述 {len(descs)} 条")
                unmatched, seen = l_annotate(str(input_pdf), str(out_pdf),
                                             lookup, panel_sys, panel_desc, descs)
                schedule_not_seen = sorted(set(lookup.keys()) - seen)
                excel_path = out_root / f"{ts}_未匹配电缆清单.xlsx"
                n1, n2, _excel = l_write_unmatched_excel(
                    str(excel_path), unmatched, schedule_not_seen, lookup, descs)
                excel_path = Path(_excel)
                stats = {"lookup_cables": len(lookup), "panels": len(panel_sys),
                         "drawing_unmatched": n1, "schedule_not_seen": n2, "mode": "lighting"}
                meta_extra = {"excel": excel_path.name, "excel_path": str(excel_path)}
            else:
                cable_lu, equip_lu, sysdesc = s_build_lookup(schedule, sheet)
                print(f"读取电缆清册: {schedule}（sheet={sheet}）")
                print(f"  电缆查找表 {len(cable_lu)} 条；设备标签 {len(equip_lu)} 个；系统描述 {len(sysdesc)} 条")
                unmatched, recognized = s_annotate(str(input_pdf), str(out_pdf),
                                                   cable_lu, equip_lu, sysdesc)
                excel_path = out_root / f"{ts}_未匹配电缆清单.xlsx"
                if unmatched:
                    s_write_unmatched_excel(str(excel_path), unmatched)
                backfill = out_root / f"{ts}_清册_含SLD识别.xlsx"
                rec_rows = [(s_drawing_no_from_path(input_pdf), pg, cab, bt, bd, et, ed, st)
                            for (pg, cab, bt, bd, et, ed, st) in recognized]
                if rec_rows:
                    backfill = Path(s_append_recognized_to_schedule(schedule, str(backfill), rec_rows))
                stats = {"cables": len(cable_lu), "equip_tags": len(equip_lu),
                         "unmatched": len(unmatched), "recognized": len(recognized),
                         "backfill_rows": len(rec_rows), "mode": "sld"}
                meta_extra = {"excel": excel_path.name, "excel_path": str(excel_path),
                              "backfill": backfill.name, "backfill_path": str(backfill)}
        except Exception as e:  # noqa: BLE001
            return {"type": "text", "content": f"标注失败：{e}"}

        main_file = out_pdf.name
        meta = {
            "path": str(out_pdf),
            "filename": main_file,
            "stats": stats,
        }
        meta.update(meta_extra)
        return {"type": "file", "content": main_file, "meta": meta}
